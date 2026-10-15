#!/usr/bin/env python3
"""Builds the submission formats of the report from the Markdown master document.

Input : docs/Documentation.md         (the master document, plain Markdown)
Output: docs/SCO310-QR-Component-Suite-Report.html
        docs/SCO310-QR-Component-Suite-Report.pdf    (Chromium print-to-PDF)
        docs/SCO310-QR-Component-Suite-Report.docx   (python-docx)

Both rendered formats gain an "Appendix D - Screenshot gallery" that embeds every figure and
every screenshot, so the submitted document is self-contained: a marker can read the report and
see the code listings and the running programs without opening a second folder.

Usage: python3 tools/make_documentation.py
"""
from __future__ import annotations

import glob
import html as html_module
import os
import re
import shutil
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, "docs")
MASTER = os.path.join(DOCS, "Documentation.md")
BASENAME = "SCO310-QR-Component-Suite-Report"

GALLERY_TITLE = "Appendix D - Screenshot gallery (embedded)"

# Captions for the generated gallery, keyed by file stem.
CAPTIONS = {
    "figure-1-component-model": ("Figure 1.", "The component model: roles, provided and required interfaces, and the facade as the composition point."),
    "figure-2-connection-model": ("Figure 2.", "The JavaBeans connection model applied: event source, listener, event object, delegation, bound properties and interface-typed connections."),
    "figure-3-deployment-model": ("Figure 3.", "The deployment model D = (C, E, P, L): component set, execution environment, descriptor parameters, deployment location."),
    "figure-4-encoding-pipeline": ("Figure 4.", "The five ISO/IEC 18004 encoding stages and the class that implements each one."),
    "figure-5-verification-strategy": ("Figure 5.", "Verification strategy: every group of checks anchored to an independent oracle or to a reconstruction."),
    "figure-6-artifact-gallery": ("Figure 6.", "Generated artifacts: PNG, SVG and ASCII renderings of the same symbols with their recovered payloads."),
    "01-code-provided-interface": ("Listing 1.", "api/QrCodeEncoder.java - the provided interface every encoder component implements."),
    "02-code-reed-solomon": ("Listing 2.", "engine/ReedSolomon.java - shift-register polynomial division over GF(2^8)."),
    "16-code-encoder-bean": ("Listing 3.", "components/PureJavaQrEncoderBean.java - a complete JavaBean with all property kinds."),
    "17-code-png-renderer": ("Listing 4.", "render/PngRendererBean.java - bound properties, validation and run-length drawing."),
    "18-code-history-bean": ("Listing 5.", "store/QrHistoryBean.java - the indexed property and the constrained maxEntries property."),
    "19-code-facade": ("Listing 6.", "facade/QrStudioFacadeBean.java - composition, event firing and timing."),
    "20-code-registry": ("Listing 7.", "facade/ComponentRegistry.java - catalogue, select, instantiate, configure, validate."),
    "21-code-beaninfo": ("Listing 8.", "render/PngRendererBeanBeanInfo.java - feature descriptors, expert flag, custom property editor."),
    "22-code-swing-assembly": ("Listing 9.", "ui/QrStudioFrame.java - the desktop front end listening to bound properties."),
    "23-code-cli-assembly": ("Listing 10.", "cli/QrCliAssembly.java - the command-line front end."),
    "24-code-zxing-adapter": ("Listing 11.", "components/ZxingQrEncoderAdapterBean.java - third-party composition behind our own interface."),
    "25-code-verification": ("Listing 12.", "test/EngineReferenceTest.java - the independent-oracle checks."),
    "03-gui-png-assembly": ("Screenshot 1.", "The desktop assembly, PNG renderer component: property sheet, live preview, indexed history."),
    "04-gui-svg-assembly": ("Screenshot 2.", "The same running application after the renderer component was replaced by renderer.svg.v1."),
    "05-gui-ascii-assembly": ("Screenshot 3.", "The same running application with renderer.ascii.v1: a textual artifact in the preview panel."),
    "06-cli-help": ("Screenshot 4.", "The command-line assembly: complete option surface."),
    "07-cli-component-catalogue": ("Screenshot 5.", "The component catalogue with roles, provided interfaces, required interfaces and state."),
    "08-cli-png-run": ("Screenshot 6.", "A full generation run: deployment report, result, deployment XML reuse."),
    "09-cli-ascii-symbol": ("Screenshot 7.", "A decodable symbol printed in a terminal by the ASCII renderer component."),
    "10-beaninspector-png-renderer": ("Screenshot 8.", "Run-time introspection of the PNG renderer bean."),
    "11-beaninspector-history-facade": ("Screenshot 9.", "Introspection of the indexed/constrained history bean and the facade event set."),
    "12-verification-transcript": ("Screenshot 10.", "The verification transcript: 108 checks, 108 passed, 0 failed."),
    "13-build-and-package": ("Screenshot 11.", "Build, verify and package in one command."),
    "14-jar-manifest": ("Screenshot 12.", "The deployment interface: Java-Bean: True entries for all nine beans."),
    "15-http-service": ("Screenshot 13.", "The HTTP assembly answering /qr, /stats and /history."),
    "26-bean-feature-matrix": ("Screenshot 14.", "The bean feature matrix across four components in one view."),
    "http-service-symbol": ("Screenshot 15.", "The artifact returned by the HTTP service for a PNG request."),
}


def read_master() -> str:
    with open(MASTER, encoding="utf-8") as handle:
        return handle.read()


def gallery_items() -> list[tuple[str, str, str]]:
    """@return (path, label, caption) for every embedded image, figures first."""
    items = []
    for path in sorted(glob.glob(os.path.join(DOCS, "diagrams", "figure-*.png"))):
        stem = os.path.basename(path)[:-4]
        label, caption = CAPTIONS.get(stem, ("Figure.", stem))
        items.append((path, label, caption))
    for path in sorted(glob.glob(os.path.join(DOCS, "screenshots", "*.png"))):
        stem = os.path.basename(path)[:-4]
        label, caption = CAPTIONS.get(stem, ("Screenshot.", stem))
        items.append((path, label, caption))
    return items


def appendix_markdown() -> str:
    lines = [f"\n---\n\n## {GALLERY_TITLE}\n",
             "The images below are the figures and the screenshots referenced throughout this "
             "report. All of them were generated from the running code by the project's own "
             "tooling (`tools/make_diagrams.py`, `tools/capture_screenshots.py`).\n"]
    for path, label, caption in gallery_items():
        rel = os.path.relpath(path, DOCS)
        lines.append(f"**{label}** {caption}\n")
        lines.append(f"![{label}]({rel})\n")
    return "\n".join(lines)


CSS = """
@page { size: A4; margin: 16mm 14mm; }
html { font-size: 11pt; }
body { font-family: "Liberation Serif", Georgia, serif; line-height: 1.5; color: #16181d;
       max-width: 200mm; margin: 0 auto; padding: 0 6mm; }
h1 { font-family: "Liberation Sans", Arial, sans-serif; font-size: 22pt; line-height: 1.2;
     border-bottom: 3px solid #1b5e20; padding-bottom: 6px; margin-top: 28px; }
h1:first-child { border-bottom: none; font-size: 26pt; margin-top: 0; }
h2 { font-family: "Liberation Sans", Arial, sans-serif; font-size: 15pt; line-height: 1.3;
     color: #1b5e20; margin-top: 26px; border-bottom: 1px solid #cfd8dc; padding-bottom: 3px; }
h3 { font-family: "Liberation Sans", Arial, sans-serif; font-size: 12.5pt; margin-top: 18px; }
code { font-family: "Liberation Mono", "DejaVu Sans Mono", monospace; font-size: 9pt;
       background: #f1f3f5; padding: 1px 3px; border-radius: 3px; }
pre { background: #16181d; color: #e6e6e6; padding: 10px 12px; border-radius: 5px;
      white-space: pre-wrap; word-break: break-word; page-break-inside: avoid; }
pre code { background: none; color: inherit; padding: 0; font-size: 8.6pt; line-height: 1.35; }
table { border-collapse: collapse; width: 100%; margin: 10px 0 16px; font-size: 9.4pt;
        page-break-inside: avoid; }
th { background: #1b5e20; color: #fff; text-align: left; padding: 5px 7px;
     font-family: "Liberation Sans", Arial, sans-serif; }
td { border: 1px solid #d7dde2; padding: 4px 7px; vertical-align: top; }
tr:nth-child(even) td { background: #f7f9fa; }
blockquote { border-left: 4px solid #90a4ae; margin: 12px 0; padding: 2px 14px; color: #37474f;
             background: #f7f9fa; }
hr { border: none; border-top: 1px solid #cfd8dc; margin: 22px 0; }
img { max-width: 100%; border: 1px solid #cfd8dc; margin: 6px 0 2px; page-break-inside: avoid; }
a { color: #1b5e20; }
p { margin: 7px 0; }
"""


def build_html() -> str:
    try:
        import markdown
    except ImportError:
        sys.exit("python-markdown is required: pip install markdown")
    body = markdown.markdown(read_master() + appendix_markdown(),
                             extensions=["tables", "fenced_code", "sane_lists", "attr_list"])
    page = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>SCO 310 Component Programming - QR Code Generator from JavaBeans Components</title>
<style>{CSS}</style>
</head>
<body>
{body}
</body>
</html>
"""
    target = os.path.join(DOCS, BASENAME + ".html")
    with open(target, "w", encoding="utf-8") as handle:
        handle.write(page)
    print("wrote", os.path.relpath(target, ROOT), f"({len(page)} chars)")
    return target


def build_pdf(html_path: str) -> None:
    browser = shutil.which("chromium") or shutil.which("chromium-browser") or shutil.which("google-chrome")
    target = os.path.join(DOCS, BASENAME + ".pdf")
    if browser is None:
        print("no Chromium/Chrome found - skipping the PDF (the HTML prints identically)")
        return
    command = [browser, "--headless", "--disable-gpu", "--no-sandbox", "--no-pdf-header-footer",
               "--virtual-time-budget=6000", f"--print-to-pdf={target}", "file://" + html_path]
    result = subprocess.run(command, capture_output=True, text=True, timeout=600)
    if os.path.exists(target):
        print("wrote", os.path.relpath(target, ROOT), f"({os.path.getsize(target)} bytes)")
    else:
        print("PDF generation failed:", result.stdout[-400:], result.stderr[-400:])


INLINE = [
    (re.compile(r"\*\*(.+?)\*\*"), r"\1", "bold"),
    (re.compile(r"(?<!\*)\*([^*]+)\*(?!\*)"), r"\1", "italic"),
    (re.compile(r"`([^`]+)`"), r"\1", "code"),
    (re.compile(r"\[([^\]]+)\]\([^)]*\)"), r"\1", None),
]


def strip_inline(text: str) -> str:
    for pattern, replacement, _ in INLINE:
        text = pattern.sub(replacement, text)
    return text


def add_runs(paragraph, text: str) -> None:
    """Adds text to a python-docx paragraph, honouring **bold**, *italic* and `code`."""
    tokens = re.split(r"(\*\*.+?\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = None
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            paragraph.add_run(token[1:-1]).italic = True
        else:
            paragraph.add_run(token)


def build_docx() -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
        from PIL import Image
    except ImportError as error:
        print("python-docx/Pillow unavailable - skipping DOCX:", error)
        return

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    lines = (read_master() + appendix_markdown()).splitlines()
    index = 0
    in_code = False
    code_buffer: list[str] = []

    def flush_code() -> None:
        nonlocal code_buffer
        if not code_buffer:
            return
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(8)
        run = paragraph.add_run("\n".join(code_buffer))
        run.font.name = "Consolas"
        run.font.size = Pt(8)
        paragraph.style = document.styles["Normal"]
        code_buffer = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                flush_code()
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_buffer.append(line.rstrip())
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        if re.fullmatch(r"-{3,}", stripped):
            index += 1
            continue

        image = re.fullmatch(r"!\[[^\]]*\]\(([^)]+)\)", stripped)
        if image:
            path = os.path.join(DOCS, image.group(1))
            if os.path.exists(path):
                with Image.open(path) as handle:
                    pixel_width, pixel_height = handle.size
                # Fit the image to the text column; very tall screenshots (the feature matrix,
                # the version ladder) are additionally limited by the printable page height.
                width_inches = 6.3
                height_inches = width_inches * pixel_height / pixel_width
                max_height = 8.6
                if height_inches > max_height:
                    width_inches *= max_height / height_inches
                document.add_picture(path, width=Inches(width_inches))
            index += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = strip_inline(heading.group(2))
            if level == 1:
                paragraph = document.add_heading(text, level=0)
            else:
                paragraph = document.add_heading(text, level=min(level - 1, 4))
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)
            index += 1
            continue

        # tables
        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[index + 1].strip()):
            header = [cell.strip() for cell in stripped.strip("|").split("|")]
            rows = []
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            table = document.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for column, cell_text in enumerate(header):
                cell = table.rows[0].cells[column]
                cell.text = ""
                add_runs(cell.paragraphs[0], cell_text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in rows:
                cells = table.add_row().cells
                for column in range(min(len(header), len(row))):
                    cells[column].text = ""
                    add_runs(cells[column].paragraphs[0], row[column])
            document.add_paragraph()
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        quote = re.match(r"^>\s?(.*)$", stripped)
        if bullet:
            add_runs(document.add_paragraph(style="List Bullet"), bullet.group(1))
        elif numbered:
            add_runs(document.add_paragraph(style="List Number"), numbered.group(2))
        elif quote:
            add_runs(document.add_paragraph(style="Intense Quote"), quote.group(1))
        else:
            add_runs(document.add_paragraph(), stripped)
        index += 1

    flush_code()
    target = os.path.join(DOCS, BASENAME + ".docx")
    document.save(target)
    print("wrote", os.path.relpath(target, ROOT), f"({os.path.getsize(target)} bytes)")


def main() -> None:
    if not os.path.exists(MASTER):
        sys.exit(f"missing master document: {MASTER}")
    html_path = build_html()
    build_pdf(html_path)
    build_docx()


if __name__ == "__main__":
    main()
